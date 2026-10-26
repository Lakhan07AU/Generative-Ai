"""Build a simple overview PPT: problem, solution, architecture, working."""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import nsmap
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from lxml import etree

NAVY = RGBColor(0x0F, 0x2C, 0x59)
ACCENT = RGBColor(0x1D, 0x6B, 0xC9)
SLATE = RGBColor(0x33, 0x41, 0x55)
MUTED = RGBColor(0x64, 0x74, 0x8B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BG = RGBColor(0xF8, 0xFA, 0xFC)
CARD = RGBColor(0xEE, 0xF4, 0xFB)
LINE = RGBColor(0xCB, 0xD5, 0xE1)


def set_run(run, text, size=18, bold=False, color=SLATE, font="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_rect(slide, l, t, w, h, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def add_round(slide, l, t, w, h, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    # Softer corners
    try:
        shape.adjustments[0] = 0.08
    except Exception:
        pass
    return shape


def tb(shape, text, size=16, bold=False, color=SLATE, align=PP_ALIGN.LEFT):
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    p.clear()
    run = p.add_run()
    set_run(run, text, size=size, bold=bold, color=color)
    return tf


def add_text(slide, l, t, w, h, text, size=16, bold=False, color=SLATE, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(l, t, w, h)
    tb(box, text, size=size, bold=bold, color=color, align=align)
    return box


def bullets(slide, l, t, w, h, items, size=18):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.level = 0
        p.space_after = Pt(10)
        run = p.add_run()
        set_run(run, "•  " + item, size=size, color=SLATE)
    return box


def footer(slide, n, total):
    add_text(slide, Inches(0.5), Inches(7.15), Inches(8), Inches(0.25),
             "AI Forensic Investigation System  ·  academic prototype", size=11, color=MUTED)
    add_text(slide, Inches(11.4), Inches(7.15), Inches(1.3), Inches(0.25),
             f"{n} / {total}", size=11, color=MUTED, align=PP_ALIGN.RIGHT)


def header_bar(slide, title, subtitle=None):
    add_rect(slide, Inches(0), Inches(0), Inches(13.333), Inches(1.15), NAVY)
    add_rect(slide, Inches(0), Inches(1.15), Inches(13.333), Inches(0.08), ACCENT)
    add_text(slide, Inches(0.5), Inches(0.22), Inches(12), Inches(0.45),
             title, size=28, bold=True, color=WHITE)
    if subtitle:
        add_text(slide, Inches(0.5), Inches(0.68), Inches(12), Inches(0.32),
                 subtitle, size=14, color=RGBColor(0xBF, 0xDB, 0xFE))


def card(slide, l, t, w, h, title, body, title_size=15, body_size=13):
    shp = add_round(slide, l, t, w, h, CARD)
    add_text(slide, l + Inches(0.18), t + Inches(0.12), w - Inches(0.3), Inches(0.35),
             title, size=title_size, bold=True, color=NAVY)
    box = slide.shapes.add_textbox(l + Inches(0.18), t + Inches(0.48), w - Inches(0.36), h - Inches(0.6))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    set_run(run, body, size=body_size, color=SLATE)
    return shp


def blank_bg(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, Inches(0), Inches(0), prs.slide_width, prs.slide_height, BG)
    return slide


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    total = 7

    # 1 Title
    s = blank_bg(prs)
    add_rect(s, Inches(0), Inches(0), Inches(13.333), Inches(7.5), NAVY)
    add_rect(s, Inches(0), Inches(0), Inches(0.18), Inches(7.5), ACCENT)
    add_text(s, Inches(0.9), Inches(1.8), Inches(11.5), Inches(0.4),
             "ACADEMIC PROTOTYPE", size=14, bold=True, color=ACCENT)
    add_text(s, Inches(0.9), Inches(2.25), Inches(11.5), Inches(1.3),
             "AI Forensic Investigation System", size=40, bold=True, color=WHITE)
    add_text(s, Inches(0.9), Inches(3.6), Inches(11), Inches(0.8),
             "Natural-language investigation over CCTV and live cameras —\nwith evidence you can trace, not guesses you cannot.",
             size=20, color=RGBColor(0xDB, 0xEA, 0xFE))
    add_text(s, Inches(0.9), Inches(5.5), Inches(11), Inches(0.4),
             "Problem  ·  Solution  ·  Architecture  ·  Working",
             size=16, color=RGBColor(0x93, 0xC5, 0xFD))

    # 2 Problem
    s = blank_bg(prs)
    header_bar(s, "Problem statement", "Why this project exists")
    bullets(s, Inches(0.55), Inches(1.5), Inches(12.2), Inches(2.4), [
        "Investigators still watch hours of CCTV by hand to answer a simple question.",
        "Clips, cameras, and timestamps sit in separate tools, so the story is easy to lose.",
        "Generic AI chat can invent people, times, or intent — that is not usable as evidence.",
        "Live cameras and uploaded footage are treated as two different worlds.",
    ], size=20)
    note = add_round(s, Inches(0.55), Inches(5.15), Inches(12.2), Inches(1.55), RGBColor(0xFF, 0xED, 0xD5))
    add_text(s, Inches(0.75), Inches(5.3), Inches(11.8), Inches(1.2),
             "Need: an assistant that finds relevant footage in plain language, keeps original evidence immutable, never claims identity or invents timestamps, and leaves the final decision with a human reviewer.",
             size=16, color=RGBColor(0x9A, 0x34, 0x12))
    footer(s, 2, total)

    # 3 Solution
    s = blank_bg(prs)
    header_bar(s, "Solution", "What we built")
    add_text(s, Inches(0.55), Inches(1.45), Inches(12.2), Inches(0.55),
             "A GenAI investigation assistant for security staff — not an autonomous security system.",
             size=18, color=SLATE)
    items = [
        ("Upload CCTV", "FFmpeg splits video into clips. YOLO detects objects. Tracks get IDs such as Person-001 — never a name or face ID."),
        ("Watch live", "Webcam, USB, RTSP, IP, phone (WebRTC), or demo file. Same pipeline: detect → track → optional VLM → evidence."),
        ("Ask in English", "Search clips and live evidence with a natural-language query, scoped to the cameras in the case."),
        ("Review and report", "A bounded agent proposes findings. A human approves or rejects. Then a forensic PDF is generated."),
    ]
    for i, (title, body) in enumerate(items):
        col = i % 2
        row = i // 2
        l = Inches(0.5) + col * Inches(6.4)
        t = Inches(2.15) + row * Inches(2.2)
        card(s, l, t, Inches(6.15), Inches(2.0), title, body, title_size=18, body_size=14)
    footer(s, 3, total)

    # 4 Architecture
    s = blank_bg(prs)
    header_bar(s, "Architecture", "Layers and stores")
    layers = [
        ("UI", "Next.js 14  ·  login, videos, live, search, cases, reports"),
        ("API", "FastAPI  ·  REST + WebSockets  ·  JWT roles"),
        ("Ingest", "Upload pipeline  and  live CameraSource session"),
        ("Perception", "YOLO  ·  tracking  ·  optional VLM / Whisper"),
        ("Evidence", "Postgres ledger  ·  MinIO bytes + SHA-256  ·  Qdrant index"),
        ("Reasoning", "RAG  ·  LangGraph agent  ·  forensic timeline + PDF"),
    ]
    for i, (name, desc) in enumerate(layers):
        y = Inches(1.45) + i * Inches(0.82)
        add_round(s, Inches(0.5), y, Inches(2.2), Inches(0.7), NAVY)
        add_text(s, Inches(0.55), y + Inches(0.18), Inches(2.1), Inches(0.4),
                 name, size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_round(s, Inches(2.9), y, Inches(9.9), Inches(0.7), CARD)
        add_text(s, Inches(3.1), y + Inches(0.18), Inches(9.5), Inches(0.4),
                 desc, size=16, color=SLATE)
    footer(s, 4, total)

    # 5 Working: ingest
    s = blank_bg(prs)
    header_bar(s, "Working — how footage becomes evidence", "Two paths, one ledger")
    card(s, Inches(0.45), Inches(1.5), Inches(6.1), Inches(5.2),
         "Path A  ·  uploaded CCTV",
         "1. Investigator uploads MP4/MOV (linked to a camera).\n"
         "2. Original file is stored and treated as immutable.\n"
         "3. FFmpeg reads metadata and extracts clips.\n"
         "4. Scene detection cuts the video at real cuts.\n"
         "5. YOLO draws boxes; tracker assigns Person-001 IDs.\n"
         "6. Clips, detections, and events are saved.\n"
         "7. Content is indexed so it can be searched later.",
         title_size=18, body_size=15)
    card(s, Inches(6.8), Inches(1.5), Inches(6.1), Inches(5.2),
         "Path B  ·  live camera",
         "1. Start a session (one active session per camera).\n"
         "2. Frames enter: ingest → sample → rolling buffer.\n"
         "3. YOLO runs on sampled frames.\n"
         "4. Tracker emits enter / leave / motion events.\n"
         "5. VLM may describe selected frames (or skip).\n"
         "6. A small JPEG is stored with SHA-256.\n"
         "7. Postgres + MinIO + Qdrant are updated together.",
         title_size=18, body_size=15)
    footer(s, 5, total)

    # 6 Working: investigation
    s = blank_bg(prs)
    header_bar(s, "Working — how an investigation runs", "Ask → retrieve → review → report")
    steps = [
        ("1", "Ask", "Type a question, e.g. “person with a backpack near camera 2 after 18:00”."),
        ("2", "Retrieve", "RAG searches Qdrant, then filters by camera, time, and objects. Out-of-case cameras are excluded."),
        ("3", "Agent", "A capped LangGraph run: classify, plan, retrieve, analyze, verify. It stops if the budget is exceeded."),
        ("4", "Review", "If findings are reviewable, status is READY FOR REVIEW. A human approves, rejects, or cancels."),
        ("5", "Forensics", "Only evidence the run already used is correlated into a timeline. Gaps and conflicts are listed, not hidden."),
        ("6", "Report", "A PDF is written with claims, evidence IDs, and verification labels: Observed / Inferred / Unknown / Conflicting."),
    ]
    for i, (num, title, body) in enumerate(steps):
        col = i % 3
        row = i // 3
        l = Inches(0.45) + col * Inches(4.25)
        t = Inches(1.5) + row * Inches(2.55)
        add_round(s, l, t, Inches(4.05), Inches(2.35), CARD)
        badge = add_round(s, l + Inches(0.18), t + Inches(0.18), Inches(0.42), Inches(0.42), ACCENT)
        add_text(s, l + Inches(0.18), t + Inches(0.22), Inches(0.42), Inches(0.35),
                 num, size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_text(s, l + Inches(0.7), t + Inches(0.22), Inches(3.1), Inches(0.4),
                 title, size=18, bold=True, color=NAVY)
        add_text(s, l + Inches(0.2), t + Inches(0.75), Inches(3.65), Inches(1.4),
                 body, size=14, color=SLATE)
    footer(s, 6, total)

    # 7 Rules / close
    s = blank_bg(prs)
    header_bar(s, "What the system will not do", "Safety rules that shape the working")
    rules = [
        ("No face recognition", "No names, biometrics, or identity claims. Only visual track IDs."),
        ("No invented time", "Timestamps come from video metadata or the live clock — never from the model."),
        ("No fake AI results", "If YOLO or the LLM is missing, the pipeline skips or answers UNKNOWN."),
        ("Human owns the close", "The agent assists. A reviewer must approve before a forensic report is final."),
    ]
    for i, (title, body) in enumerate(rules):
        y = Inches(1.5) + i * Inches(1.2)
        add_round(s, Inches(0.5), y, Inches(12.3), Inches(1.05), CARD)
        add_text(s, Inches(0.75), y + Inches(0.12), Inches(11.8), Inches(0.35),
                 title, size=18, bold=True, color=NAVY)
        add_text(s, Inches(0.75), y + Inches(0.5), Inches(11.8), Inches(0.4),
                 body, size=16, color=SLATE)
    footer(s, 7, total)

    out = r"D:\Github\Generative-Ai\AI Forensic Investigation System\Document\AI-Forensic-Investigation-Overview.pptx"
    prs.save(out)
    print(out)


if __name__ == "__main__":
    build()
